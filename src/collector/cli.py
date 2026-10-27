"""コマンドライン入口。

  python -m collector run [--storage local|release ...] [--region KEY ...] [--dry-run] [--report-json PATH]
  python -m collector inspect FILE ...

終了コード: 0=成功 / 1=いずれかの地域が失敗 / 2=設定不備などで処理全体が止まった
"""

from __future__ import annotations

import argparse
import io
import json
import logging
import os
import sys
import tempfile
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

from . import notify
from .archive import EXCEL_EXTS, open_zip
from .config import Config, load_config
from .http import PoliteClient
from .runner import run
from .storage import GitHubReleaseStorage, LocalStorage, Storage

log = logging.getLogger("collector")


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", stream=sys.stderr)
    args = _parser().parse_args(argv)
    if args.command == "inspect":
        return inspect_files([Path(p) for p in args.paths])
    return _run(args)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="collector", description="厚生局3地域の保険薬局指定一覧を収集する")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("run", help="収集を実行する")
    p.add_argument("--config", default="config.yaml")
    p.add_argument("--manifest", default="state/manifest.json")
    p.add_argument("--storage", action="append", choices=["local", "release"],
                   help="保存先（複数指定可）。省略時は local")
    p.add_argument("--data-dir", default="data", help="local 保存先のフォルダ")
    p.add_argument("--region", action="append", help="対象の地域キー（複数指定可）。省略時は全地域")
    p.add_argument("--dry-run", action="store_true", help="一覧ページの抽出結果だけ確認する（ダウンロード・保存・通知なし）")
    p.add_argument("--report-json", help="実行結果をJSONで書き出すパス")

    p = sub.add_parser("inspect", help="ダウンロード済みのZIP/Excelの中身（ファイル名・シート）を表示する")
    p.add_argument("paths", nargs="+")
    return parser


def _run(args: argparse.Namespace) -> int:
    now = datetime.now(timezone.utc)
    webhook = os.environ.get("SLACK_WEBHOOK_URL") if not args.dry_run else None
    run_url = _github_run_url()
    try:
        config = load_config(args.config)
        storages = _make_storages(args.storage or ["local"], args.data_dir, config)
        with tempfile.TemporaryDirectory(prefix="collector-") as staging:
            results = run(config, PoliteClient(config.http), args.manifest, storages, Path(staging),
                          regions=args.region, dry_run=args.dry_run, now=now)
    except Exception as e:
        log.exception("収集処理が異常終了しました")
        if webhook:
            text = f":x: *厚生局 保険薬局指定一覧の収集処理が異常終了しました*\n{type(e).__name__}: {e}"
            _post(webhook, text + (f"\n<{run_url}|実行ログを開く>" if run_url else ""))
        return 2

    print(notify.format_summary(results))
    if args.report_json:
        report = {"started_at": now.isoformat(), "dry_run": args.dry_run, "results": [asdict(r) for r in results]}
        Path(args.report_json).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if webhook and notify.needs_notification(results):
        _post(webhook, notify.build_message(results, now, run_url))
    return 1 if any(r.status == "failed" for r in results) else 0


def _make_storages(names: list[str], data_dir: str, config: Config) -> list[Storage]:
    storages: list[Storage] = []
    for name in dict.fromkeys(names):
        if name == "local":
            storages.append(LocalStorage(data_dir))
        elif name == "release":
            storages.append(GitHubReleaseStorage(notes=_release_notes(config)))
    return storages


def _release_notes(config: Config) -> str:
    sources = "\n".join(f"- {r.label}ホームページ（{r.page_url}）" for r in config.regions.values())
    return (
        "各地方厚生局のホームページで公開されている「コード内容別医療機関一覧表（薬局）」の原本です。\n"
        "利用条件: 公共データ利用規約（第1.0版）。加工して利用する場合は、加工した旨を記載してください。\n\n"
        f"出典:\n{sources}\n\n取得日時・SHA256 はリポジトリの state/manifest.json を参照してください。"
    )


def _github_run_url() -> str | None:
    server, repo, run_id = (os.environ.get(k) for k in ("GITHUB_SERVER_URL", "GITHUB_REPOSITORY", "GITHUB_RUN_ID"))
    return f"{server}/{repo}/actions/runs/{run_id}" if server and repo and run_id else None


def _post(webhook: str, text: str) -> None:
    try:
        notify.post_slack(webhook, text)
    except Exception as e:   # 通知の失敗で収集結果（終了コード）を変えない
        log.error("Slack通知に失敗しました: %s", e)


def inspect_files(paths: list[Path]) -> int:
    for path in paths:
        print(f"== {path}")
        data = path.read_bytes()
        suffix = path.suffix.lower()
        if suffix == ".zip":
            with open_zip(data) as zf:
                for info in zf.infolist():
                    print(f"  {info.filename}  {info.file_size:,} バイト")
                    if Path(info.filename).suffix.lower() in EXCEL_EXTS:
                        _print_sheets(zf.read(info), Path(info.filename).suffix.lower(), indent="    ")
        elif suffix in EXCEL_EXTS:
            _print_sheets(data, suffix, indent="  ")
        else:
            print("  （ZIP/Excel以外は表示できません）")
    return 0


def _print_sheets(data: bytes, suffix: str, indent: str) -> None:
    if suffix == ".xls":
        print(f"{indent}旧形式（.xls）のためシートは表示できません")
        return
    try:
        import openpyxl
    except ImportError:
        print(f"{indent}openpyxl が未インストールのためシートは表示できません（pip install '.[inspect]'）")
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    for ws in wb.worksheets:
        print(f"{indent}シート「{ws.title}」 {ws.max_row}行 × {ws.max_column}列")
    wb.close()
